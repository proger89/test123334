from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
import re
root=Path(__file__).resolve().parents[1]
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.75)
sec.left_margin=sec.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1','Heading 2']:
 style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(8)
doc.core_properties.author='Команда НаракаТОПчик'
doc.core_properties.title='Виртуальная смена ВСМ Техническое задание версии 4.1'
for line in (root/'docs/VSM_TZ_v4_1.md').read_text(encoding='utf8').splitlines():
 if not line.strip():continue
 text=re.sub(r'`([^`]+)`',r'\1',line)
 if text.startswith('# '):doc.add_paragraph(text[2:],'Title')
 elif text.startswith('## '):doc.add_heading(text[3:],level=1)
 else:doc.add_paragraph(text)
doc.save(root/'docs/VSM_TZ_v4_1.docx')
